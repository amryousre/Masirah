"""Text from native PDF and DOCX CVs (T052, FR-004).

Extraction reads whatever this module returns, so a value missing here cannot be extracted
correctly later. The corpus is the fixture: every value labelled for a native PDF, Arabic
included, must be in its text, in the right order.
"""

import io
import json
import re
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

import pymupdf
import pytest
from docx import Document
from docx.oxml import parse_xml

from masirah_pipeline.adapters.native_text import read_native_text, readable
from masirah_pipeline.adapters.visual_order import brackets_balanced, logical_line
from masirah_pipeline.ports.document_text import UnreadableDocumentError

CORPUS = Path(__file__).resolve().parents[1] / "fixtures" / "cvs"


def labels(path: Path) -> dict[str, Any]:
    return json.loads(path.with_suffix(".json").read_text(encoding="utf-8"))  # type: ignore[no-any-return]


NATIVE = [p for p in sorted(CORPUS.glob("*.pdf")) if labels(p)["document"]["capture"] == "native"]
CAPTURED = [p for p in sorted(CORPUS.glob("*.pdf")) if labels(p)["document"]["capture"] != "native"]

ARABIC_INDIC = {base + n: str(n) for base in (0x0660, 0x06F0) for n in range(10)}
# Letter forms a CV may write either way, and the marks it may or may not write.
FOLDED = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه"})


def letters(text: str) -> str:
    """Letters and digits only, folded, so only the characters and their order are compared."""
    text = unicodedata.normalize("NFKC", text).casefold().translate(ARABIC_INDIC).translate(FOLDED)
    return "".join(char for char in text if char.isalnum())


def label_values(profile: dict[str, Any]) -> list[str]:
    values = [profile["name"], profile["location"], *profile["emails"], *profile["skills"]]
    for job in profile["employment"]:
        values += [job["employer"], job["title"]]
    for school in profile["education"]:
        values += [school["institution"], school["qualification"]]
    for language in profile["languages"]:
        values += [language["language"], language["level"]]
    for cert in profile["certifications"]:
        values += [cert["name"], cert["issuer"]]
    return [value for value in values if value]


def stem(path: Path) -> str:
    return path.stem


@pytest.mark.parametrize("pdf", NATIVE, ids=stem)
def test_every_labelled_value_is_in_the_text_in_order(pdf: Path) -> None:
    """Arabic included: a value stored in visual order reads backwards until it is repaired."""
    text = letters(read_native_text(pdf.read_bytes()))
    missing = [v for v in label_values(labels(pdf)["profile"]) if letters(v) not in text]
    assert missing == []


@pytest.mark.parametrize("pdf", NATIVE, ids=stem)
def test_every_labelled_phone_and_date_is_in_the_text(pdf: Path) -> None:
    text = read_native_text(pdf.read_bytes()).translate(ARABIC_INDIC)
    profile = labels(pdf)["profile"]
    digits = re.sub(r"\D", "", text)
    assert [phone for phone in profile["phones"] if phone[-10:] not in digits] == []
    entries = profile["employment"] + profile["education"]
    years = [
        e[key][:4] for e in entries for key in ("start", "end") if e[key] not in (None, "present")
    ]
    assert [year for year in years if year not in text] == []


def test_text_inside_an_embedded_figure_is_read() -> None:
    """Some CV builders draw text inside a figure (a PDF form XObject), not on the page."""
    inner = pymupdf.open()
    inner.new_page().insert_text((72, 72), "Senior Accountant")
    outer = pymupdf.open()
    page = outer.new_page()
    page.show_pdf_page(page.rect, inner, 0)
    assert "Senior Accountant" in read_native_text(outer.tobytes())


@pytest.mark.parametrize("pdf", CAPTURED, ids=stem)
def test_a_scan_or_photo_has_no_native_text(pdf: Path) -> None:
    """Empty text is what sends a document to OCR instead."""
    assert read_native_text(pdf.read_bytes()).strip() == ""


class TestVisualOrder:
    """Right-to-left lines stored left to right as they appear, as Chromium writes them."""

    @pytest.mark.parametrize(
        ("stored", "logical"),
        [
            # Corpus lines, as pdfminer reads them from syn-ar-native-01 and syn-mixed-native-01.
            ("ديمحﻟا دبﻋ ريمﺳ دمﺣأ", "أحمد سمير عبد الحميد"),
            ("نﻵا ﻰتﺣ - 2021 سرﺎﻣ", "مارس 2021 - حتى الآن"),
            ("2021 رﯾاربﻓ - 2017 ربمتبﺳ", "سبتمبر 2017 - فبراير 2021"),
            (
                "ahmed.samir@example.com  |  0100 000 0201  |  رصﻣ ،ةرﻫﺎﻘﻟا",
                "القاهرة، مصر  |  0100 000 0201  |  ahmed.samir@example.com",
            ),
            (
                "English: Very Good  |  مﻷا ﺔﻐﻠﻟا :ﺔﯾﺑرﻌﻟا",
                "العربية: اللغة الأم  |  English: Very Good",
            ),
            ("Senior Medical Representative", "Senior Medical Representative"),
            (".target", "target."),
            (
                "(Front-End Developer) ﺔﯿﻣﺎﻣأ تﺎﮫﺟاو ةرﻮﻄﻣ",
                "مطورة واجهات أمامية (Front-End Developer)",
            ),
            # A mark read before its letter, with a space pdfminer put between them.
            ("ا ًدﺟ ديﺟ", "جيد جدًا"),
            # Arabic-Indic digit groups are laid out right to left even in a left-to-right span,
            # from syn-ar-native-02.
            ("٠٢٠٢ ٠٠٠ ٠١٠٠", "٠١٠٠ ٠٠٠ ٠٢٠٢"),
            # A year after Latin text belongs with it, from syn-mixed-native-02.
            (
                "Power BI - ITI، 2020 ماﺪﺨﺘﺳﺎﺑ تﺎﻧﺎﻴﺒﻟا ﻞﻴﻠﺤﺗ",
                "تحليل البيانات باستخدام Power BI - ITI، 2020",
            ),
        ],
    )
    def test_a_right_to_left_line_is_put_in_reading_order(self, stored: str, logical: str) -> None:
        assert readable(logical_line(stored, right_to_left=True)) == logical

    def test_a_field_after_a_separator_is_laid_out_on_its_own(self) -> None:
        """Templates set "|"-separated fields in their own spans, so a year range after one
        keeps its order rather than joining the Arabic before it."""
        stored = "ﺓﺮﻫﺎﻘﻟا | 2017 - 2020"
        assert readable(logical_line(stored, right_to_left=False)) == "القاهرة | 2017 - 2020"

    @pytest.mark.parametrize(
        ("text", "balanced"),
        [
            ("طاقم العمل ( 2006 - 2011 )", True),
            ("no brackets", True),
            ("(a [b] c)", True),
            (") 2006 - 2011 )", False),
            ("(a ] b [ c)", False),  # every depth is fine, but the pairs cross
            ("(a [b) c]", False),
            ("(open", False),
        ],
    )
    def test_brackets_must_open_and_close_in_matching_pairs(
        self, text: str, balanced: bool
    ) -> None:
        assert brackets_balanced(text) is balanced

    def test_arabic_in_a_left_to_right_line_is_reversed_in_place(self) -> None:
        stored = "Native: ﺔﯿﺑﺮﻌﻟا"
        assert readable(logical_line(stored, right_to_left=False)) == "Native: العربية"


def docx_bytes(document: Any) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def text_box(paragraph: Any, text: str) -> None:
    """A floating text box anchored in the paragraph, as Word CV templates place a sidebar."""
    namespaces = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml"'
    )
    content = f"<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
    run = parse_xml(
        f"<w:r {namespaces}><mc:AlternateContent>"
        f'<mc:Choice Requires="wps"><wps:txbx>{content}</wps:txbx></mc:Choice>'
        f"<mc:Fallback><v:textbox>{content}</v:textbox></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )
    paragraph._p.append(run)


class TestDocx:
    def test_paragraphs_tables_headers_and_text_boxes_are_read_in_order(self) -> None:
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Mona Khaled"
        document.add_paragraph("Front-end developer")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Experience\nNile Pharma 2020 - 2023"
        table.cell(0, 1).text = "Skills\nReact"
        anchor = document.add_paragraph("Education")
        text_box(anchor, "Languages: Arabic, English")
        document.add_paragraph("مطورة واجهات أمامية")
        rows = [row for row in read_native_text(docx_bytes(document)).splitlines() if row]
        assert rows == [
            "Mona Khaled",
            "Front-end developer",
            "Experience",
            "Nile Pharma 2020 - 2023",
            "Skills",
            "React",
            "Education",
            "Languages: Arabic, English",
            "مطورة واجهات أمامية",
        ]

    def test_a_text_box_is_read_once_though_word_stores_it_twice(self) -> None:
        document = Document()
        text_box(document.add_paragraph(), "Cairo")
        assert read_native_text(docx_bytes(document)).count("Cairo") == 1

    def test_each_section_is_read_with_its_own_header_and_footer(self) -> None:
        """A later section can carry its own header, and it is read where that section starts."""
        document = Document()
        first = document.sections[0]
        first.header.paragraphs[0].text = "Mona Khaled"
        first.footer.paragraphs[0].text = "Page one footer"
        document.add_paragraph("Experience")
        second = document.add_section()
        second.header.is_linked_to_previous = False
        second.header.paragraphs[0].text = "Mona Khaled, references"
        second.footer.is_linked_to_previous = False
        second.footer.paragraphs[0].text = "Page two footer"
        document.add_paragraph("References on request")
        rows = [row for row in read_native_text(docx_bytes(document)).splitlines() if row]
        assert rows == [
            "Mona Khaled",
            "Experience",
            "Page one footer",
            "Mona Khaled, references",
            "References on request",
            "Page two footer",
        ]

    def test_a_header_shared_with_the_section_before_is_read_once(self) -> None:
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Mona Khaled"
        document.add_paragraph("Experience")
        document.add_section()
        document.add_paragraph("References on request")
        assert read_native_text(docx_bytes(document)).count("Mona Khaled") == 1

    def test_a_docx_that_unpacks_beyond_the_limit_is_refused_unread(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """A small archive can unpack to gigabytes. The sizes it declares are checked first."""
        content = docx_bytes(Document())
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            unpacked = sum(entry.file_size for entry in archive.infolist())
        monkeypatch.setattr("masirah_pipeline.adapters.native_text.MAX_DOCX_BYTES", unpacked - 1)
        with pytest.raises(UnreadableDocumentError, match="unpacks"):
            read_native_text(content)


@pytest.mark.parametrize(
    "content",
    [b"", b"plain text, not a document", b"PK\x03\x04 not really a zip", b"%PDF-1.7 truncated"],
    ids=["empty", "text", "broken-zip", "broken-pdf"],
)
def test_a_file_that_is_not_a_readable_pdf_or_docx_is_refused(content: bytes) -> None:
    with pytest.raises(UnreadableDocumentError):
        read_native_text(content)


def test_a_password_protected_pdf_is_refused_as_such() -> None:
    """The review queue tells the recruiter to ask for a copy without a password (T081)."""
    locked = pymupdf.open(CORPUS / "en-native-02.pdf").tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret"
    )
    with pytest.raises(UnreadableDocumentError, match="password-protected"):
        read_native_text(locked)
